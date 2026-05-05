#!/usr/bin/env python3
"""
Generate Word (.docx) tables from GRIB2 CSV files (CodeFlag + Templates).

Usage:
    python scripts/csv2docx.py <lang> [--type codeflag|template|all] [--outdir DIR]

Examples:
    python scripts/csv2docx.py ru                          # all codeflags + templates
    python scripts/csv2docx.py ru --type codeflag          # codeflags only
    python scripts/csv2docx.py fr --type template          # French templates only
"""

import argparse, csv, os, sys, glob

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

FONT = 'Arial'
SZ_TITLE = Pt(11)
SZ_HDR = Pt(8)
SZ_DATA = Pt(8)


def load_notes(lang, base_dir, note_type):
    """Load notes for a specific type (CodeFlag or Template). No English fallback."""
    notes = {}
    if note_type == 'codeflag':
        notes_file = os.path.join(base_dir, 'notes', f'CodeFlag_notes_{lang}.csv')
    else:
        notes_file = os.path.join(base_dir, 'notes', f'Template_notes_{lang}.csv')

    if os.path.exists(notes_file):
        with open(notes_file, encoding='utf-8-sig') as f:
            for r in csv.DictReader(f):
                nid = r.get('noteID', '').strip()
                text = r.get(f'note_{lang}', r.get('note', r.get(f'note_{lang}', ''))).strip()
                if nid and text:
                    notes[nid] = text
    return notes


def resolve_note(row, lang, notes_db):
    note_text = row.get(f'Note_{lang}', '').strip()
    note_ids = row.get('noteIDs', '').strip()
    if not note_ids:
        return note_text
    resolved = []
    for nid in note_ids.split(','):
        nid = nid.strip()
        if nid in notes_db:
            resolved.append(f'Note {nid}: {notes_db[nid]}')
    return ' | '.join(resolved) if resolved else note_text


def sanitize(text):
    """Remove XML-incompatible control characters."""
    if not text:
        return ''
    return ''.join(c for c in str(text) if ord(c) >= 32 or c in '\t\n\r')


def set_cell(cell, text, size=SZ_DATA, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.space_before = Pt(1)
    p.space_after = Pt(1)
    run = p.add_run(sanitize(text))
    run.font.name = FONT
    run.font.size = size
    run.font.bold = bold


def style_header(table, headers):
    for cell, hdr in zip(table.rows[0].cells, headers):
        set_cell(cell, hdr, size=SZ_HDR, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        tcPr = cell._element.get_or_add_tcPr()
        tcPr.append(tcPr.makeelement(qn('w:shd'), {
            qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): 'D9E2F3'}))


# ── CodeFlag generator ───────────────────────────────────────────────

def generate_codeflag(lang, csv_path, doc, notes_db):
    with open(csv_path, encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    if not rows:
        return

    # Extract table info from filename (e.g., GRIB2_CodeFlag_4_2_0_0_CodeTable_ru.csv)
    basename = os.path.basename(csv_path)
    title_text = rows[0].get(f'Title_{lang}', basename)

    # Title
    p = doc.add_paragraph()
    run = p.add_run(title_text)
    run.font.name = FONT
    run.font.size = SZ_TITLE
    run.font.bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Subtitle if present
    subtitle = rows[0].get(f'SubTitle_{lang}', '').strip()
    if subtitle:
        p2 = doc.add_paragraph()
        run2 = p2.add_run(subtitle)
        run2.font.name = FONT
        run2.font.size = Pt(9)
        run2.font.italic = True
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Determine which columns to show
    has_value = any(r.get('Value', '').strip() for r in rows)
    has_units = any(r.get(f'UnitComments_{lang}', '').strip() for r in rows)

    headers = ['CodeFlag']
    if has_value:
        headers.append('Value')
    headers.append('Description')
    if has_units:
        headers.append('Units')
    headers.append('Note')

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    style_header(table, headers)

    for r in rows:
        row_cells = table.add_row().cells
        vals = [r.get('CodeFlag', '')]
        if has_value:
            vals.append(r.get('Value', ''))
        vals.append(r.get(f'MeaningParameterDescription_{lang}', ''))
        if has_units:
            vals.append(r.get(f'UnitComments_{lang}', ''))
        vals.append(resolve_note(r, lang, notes_db))
        for cell, val in zip(row_cells, vals):
            set_cell(cell, val)


# ── Template generator ───────────────────────────────────────────────

def generate_template(lang, csv_path, doc, notes_db):
    with open(csv_path, encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    if not rows:
        return

    title_text = rows[0].get(f'Title_{lang}', os.path.basename(csv_path))

    p = doc.add_paragraph()
    run = p.add_run(title_text)
    run.font.name = FONT
    run.font.size = SZ_TITLE
    run.font.bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    headers = ['Octet No.', 'Count', 'Contents', 'Note', 'Code/Flag Table']
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    style_header(table, headers)

    for r in rows:
        row_cells = table.add_row().cells
        # Combine codeTable and flagTable references
        tbl_ref = r.get('codeTable', '').strip()
        flag_ref = r.get('flagTable', '').strip()
        ref = tbl_ref or flag_ref
        if tbl_ref and flag_ref:
            ref = f'Code {tbl_ref} / Flag {flag_ref}'
        elif tbl_ref:
            ref = f'Code table {tbl_ref}'
        elif flag_ref:
            ref = f'Flag table {flag_ref}'

        vals = [
            r.get('OctetNo', ''),
            r.get('OctetCount', ''),
            r.get(f'Contents_{lang}', ''),
            resolve_note(r, lang, notes_db),
            ref,
        ]
        for cell, val in zip(row_cells, vals):
            set_cell(cell, val)


# ── Main ─────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description='Generate Word tables from GRIB2 CSVs')
    parser.add_argument('lang', help='Language code (ru, fr, es, en)')
    parser.add_argument('--type', default='all', choices=['codeflag', 'template', 'all'])
    parser.add_argument('--outdir', default=None)
    args = parser.parse_args()

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    lang_dir_map = {'en': 'english', 'fr': 'french', 'es': 'spanish', 'ru': 'russian'}
    lang_name = lang_dir_map.get(args.lang, args.lang)
    out_dir = args.outdir or os.path.join(base_dir, 'docx', args.lang)
    os.makedirs(out_dir, exist_ok=True)

    total = 0

    # CodeFlags
    if args.type in ('codeflag', 'all'):
        cf_dir = os.path.join(base_dir, lang_name, 'codeflags')
        if os.path.isdir(cf_dir):
            notes_db = load_notes(args.lang, base_dir, 'codeflag')
            print(f'CodeFlag: loaded {len(notes_db)} notes for {args.lang}')
            for f in sorted(glob.glob(os.path.join(cf_dir, '*.csv'))):
                doc = Document()
                s = doc.sections[0]
                s.orientation = WD_ORIENT.LANDSCAPE
                s.page_width = Cm(29.7)
                s.page_height = Cm(21.0)
                s.left_margin = Cm(1.5)
                s.right_margin = Cm(1.5)

                generate_codeflag(args.lang, f, doc, notes_db)
                name = os.path.splitext(os.path.basename(f))[0]
                doc.save(os.path.join(out_dir, f'{name}.docx'))
                total += 1

    # Templates
    if args.type in ('template', 'all'):
        tpl_dir = os.path.join(base_dir, lang_name, 'templates')
        if os.path.isdir(tpl_dir):
            notes_db = load_notes(args.lang, base_dir, 'template')
            print(f'Template: loaded {len(notes_db)} notes for {args.lang}')
            for f in sorted(glob.glob(os.path.join(tpl_dir, f'*_{args.lang}.csv'))):
                doc = Document()
                s = doc.sections[0]
                s.orientation = WD_ORIENT.LANDSCAPE
                s.page_width = Cm(29.7)
                s.page_height = Cm(21.0)
                s.left_margin = Cm(1.5)
                s.right_margin = Cm(1.5)

                generate_template(args.lang, f, doc, notes_db)
                name = os.path.splitext(os.path.basename(f))[0]
                doc.save(os.path.join(out_dir, f'{name}.docx'))
                total += 1

    print(f'\n{total} files written to {out_dir}')


if __name__ == '__main__':
    main()
